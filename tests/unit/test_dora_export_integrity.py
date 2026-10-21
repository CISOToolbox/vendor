"""BUG-113 — the exported register passes the key checks the EBA runs on
reception: no key reported twice in a template (rule 806) and no reference
to a key another template does not report (rule 807).

Same in-memory register as test_dora_export_complete.py."""
import io
import uuid

import pytest
import pytest_asyncio
from fastapi import HTTPException
from openpyxl import load_workbook

from src.models import Base, DoraArrangement, DoraArrangementSubcontractor, DoraSubcontractor, Vendor
from tests.unit.test_dora_export_complete import (
    LEI_FE, LEI_PARENT, _code, _engine, _export, _red_cells, _seed, _Session, _set,
)

pytestmark = pytest.mark.asyncio

# Primary key columns per template (Data Model for DORA RoI, keys in bold).
_KEYS = {
    "B_01.02": [0], "B_02.01": [0], "B_02.03": [0, 1], "B_03.01": [0, 1], "B_03.02": [0, 1],
    "B_03.03": [0, 1], "B_05.01": [0], "B_05.02": [0, 1, 2, 4, 5],
}
LEI_OUTSIDE = "969500UP76J52A9OXU27"   # a parent that is no vendor of the register


@pytest_asyncio.fixture(autouse=True)
async def _db():
    async with _engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    yield
    async with _engine.begin() as conn:
        await conn.run_sync(Base.metadata.drop_all)


def _sheets(blob: bytes):
    wb = load_workbook(io.BytesIO(blob))
    return lambda name: [[c.value for c in r] for r in wb[name].iter_rows(min_row=3)]   # noqa: E731


def _violations(blob: bytes) -> list[str]:
    rows = _sheets(blob)
    out = []
    for name, cols in _KEYS.items():
        keys = [tuple(r[c] for c in cols) for r in rows(name)]
        out += [f"806 {name} {k}" for k in {k for k in keys if keys.count(k) > 1}]
    col = lambda name, i: {r[i] for r in rows(name) if r[i] not in (None, "")}   # noqa: E731
    fks = [
        ("B_02.03", 0, "B_02.01", 0), ("B_02.03", 1, "B_02.01", 0),
        ("B_03.02", 1, "B_05.01", 0), ("B_03.03", 1, "B_01.02", 0),
        ("B_05.01", 10, "B_05.01", 0), ("B_05.02", 2, "B_05.01", 0),
    ]
    for src, i, dst, j in fks:
        out += [f"807 {src}[{i}] {v} not in {dst}" for v in sorted(col(src, i) - col(dst, j))]
    return out


async def _add(db, pid, objs):
    for o in objs(pid):
        db.add(o)


async def test_the_complete_register_passes_the_key_checks():
    assert _violations(await _export()) == []


async def test_a_parent_kept_as_a_lei_that_is_a_vendor_is_that_vendor():
    # PP-004's parent stored as PP-001's LEI: one B_05.01 row for PP-001, not two.
    blob = await _export({_set: (Vendor, ("PP-004",), "ultimate_parent_id", LEI_PARENT)})
    assert _violations(blob) == []


async def test_a_parent_kept_as_a_lei_outside_the_register_has_a_row_of_its_own():
    blob = await _export({_set: (Vendor, ("PP-004",), "ultimate_parent_id", LEI_OUTSIDE)})
    assert _violations(blob) == []
    parent = [r for r in _sheets(blob)("B_05.01") if r[0] == LEI_OUTSIDE]
    assert len(parent) == 1 and parent[0][10] == LEI_OUTSIDE   # top of its group: its own code
    assert any(c.startswith("B_05.01.0050") for c in _red_cells(blob))   # its name is still to be entered


async def test_an_intra_group_provider_outside_the_consolidation_shows_red_not_a_dangling_lei():
    # Intra-group through its parent (the reporting entity), its own LEI is in no B_01.02 row.
    def objs(pid):
        yield Vendor(project_id=pid, id="PP-005", name="Group Ops", lei="5493006MHB84DD0ZWV18",
                     ultimate_parent_id=LEI_FE, person_type=_code("person_type"), country_iso2="FR")
        yield DoraArrangement(project_id=pid, id="ARR-005", vendor_id="PP-005", arrangement_reference="ARR-005",
                              start_date="2025-01-01", end_date="2027-01-01")
    blob = await _export({_add: objs})
    assert [v for v in _violations(blob) if "B_03.03" in v] == []
    assert any(c.startswith("B_03.03.0020") for c in _red_cells(blob))


async def test_an_intra_group_arrangement_without_a_link_has_no_b02_03_row():
    rows = _sheets(await _export({_set: (DoraArrangement, ("ARR-002",), "linked_arrangement_id", None)}))
    assert [r for r in rows("B_02.03") if r[0] == "ARR-002"] == []   # B_02.03 links two arrangements
    assert [r for r in rows("B_03.03") if r[0] == "ARR-002"]          # still an intra-group provider


async def test_a_rank_change_that_breaks_the_chain_is_refused():
    from src.routes.dora import update_subcontractor_link
    from src.schemas import DoraArrangementSubcontractorUpdate
    async with _Session() as db:
        pid = uuid.uuid4()
        await _seed(db, pid)
        # SUB-001 (rank 2) is the rank n-1 of SUB-002 (rank 3): moving it to rank 3 breaks the chain.
        with pytest.raises(HTTPException) as exc:
            await update_subcontractor_link(pid, "ARR-001", "SUB-001",
                                            DoraArrangementSubcontractorUpdate(tier=3), user=None, db=db)
        assert exc.value.status_code == 422
        link = await db.get(DoraArrangementSubcontractor, (pid, "ARR-001", "SUB-001"))
        assert link.tier == 2


async def test_unlinking_a_rank_n_minus_1_clears_its_children():
    from src.routes.dora import unlink_subcontractor
    async with _Session() as db:
        pid = uuid.uuid4()
        await _seed(db, pid)
        await unlink_subcontractor(pid, "ARR-001", "SUB-001", user=None, db=db)
        child = await db.get(DoraArrangementSubcontractor, (pid, "ARR-001", "SUB-002"))
        await db.refresh(child)
        assert child.parent_subcontractor_id is None   # its B_05.02.0060 shows red, never a stale recipient
        assert await db.get(DoraSubcontractor, (pid, "SUB-001"))


async def test_a_link_is_refused_on_an_arrangement_that_is_not_intra_group():
    from src.routes.dora import update_arrangement
    from src.schemas import DoraArrangementUpdate
    async with _Session() as db:
        pid = uuid.uuid4()
        await _seed(db, pid)
        with pytest.raises(HTTPException) as exc:   # Cloud Corp is an external provider
            await update_arrangement(pid, "ARR-001", DoraArrangementUpdate(linked_arrangement_id="ARR-003"),
                                     user=None, db=db)
        assert exc.value.status_code == 422


async def test_a_subcontractor_is_saved_without_a_type_of_person_and_with_one():
    from src.routes.dora import create_subcontractor, update_subcontractor
    from src.schemas import DoraSubcontractorCreate, DoraSubcontractorUpdate
    async with _Session() as db:
        pid = uuid.uuid4()
        await _seed(db, pid)
        await create_subcontractor(pid, DoraSubcontractorCreate(id="SUB-NEW", name="New"), user=None, db=db)
        await update_subcontractor(pid, "SUB-NEW", DoraSubcontractorUpdate(person_type=""), user=None, db=db)
        await update_subcontractor(pid, "SUB-NEW", DoraSubcontractorUpdate(person_type=_code("person_type")), user=None, db=db)
        assert (await db.get(DoraSubcontractor, (pid, "SUB-NEW"))).person_type == _code("person_type")
        with pytest.raises(HTTPException) as exc:
            await update_subcontractor(pid, "SUB-NEW", DoraSubcontractorUpdate(person_type="robot"), user=None, db=db)
        assert exc.value.status_code == 422


async def test_a_lower_case_subcontractor_lei_is_the_same_key_in_both_templates():
    rows = _sheets(await _export({_set: (DoraSubcontractor, ("SUB-001",), "lei", "254900oppu84gm83mg36")}))
    assert "254900OPPU84GM83MG36" in {r[0] for r in rows("B_05.01")}
    assert "254900OPPU84GM83MG36" in {r[2] for r in rows("B_05.02")}
    assert "254900OPPU84GM83MG36" in {r[5] for r in rows("B_05.02")}   # recipient of the rank 3 one


# ── Review of BUG-113: the key checks hold whatever the data entered ──


async def _many(db, pid, changes):
    for change in changes:
        await _set(db, pid, change)


async def test_a_lower_case_vendor_lei_that_is_also_a_subcontractor_is_one_key():
    # A hyperscaler both direct provider and subcontractor; LEIs typed in lower case.
    blob = await _export({_many: [(Vendor, ("PP-002",), "lei", "213800wsgiizcxf1p572"),
                                  (DoraSubcontractor, ("SUB-002",), "lei", "213800WSGIIZCXF1P572")]})
    assert _violations(blob) == []


async def test_lower_case_entity_and_consolidation_leis_are_the_same_keys():
    from src.models import DoraConsolidationScope, DoraEntity
    blob = await _export({_many: [(DoraEntity, ("ENT-001",), "lei", LEI_FE.lower()),
                                  (DoraConsolidationScope, ("CS-001",), "entity_lei", "5493001kjtiigc8y1r12")]})
    assert _violations(blob) == []


async def test_a_provider_signer_of_its_group_is_reported_in_b05_01():
    from src.models import DoraSigner
    def objs(pid):
        yield DoraSigner(project_id=pid, arrangement_id="ARR-001", id="SIG-SUBS", signer_lei=LEI_OUTSIDE.lower(),
                         signer_name="Cloud Corp Ireland", signer_role="tpp")
    blob = await _export({_add: objs})
    assert _violations(blob) == []
    assert LEI_OUTSIDE in {r[0] for r in _sheets(blob)("B_05.01")}


async def test_two_vendors_with_one_lei_give_one_b05_01_row():
    def objs(pid):
        yield Vendor(project_id=pid, id="PP-009", name="Cloud Corp (duplicate)", lei="213800WSGIIZCXF1P572",
                     person_type=_code("person_type"), country_iso2="IE")
        yield DoraArrangement(project_id=pid, id="ARR-009", vendor_id="PP-009", arrangement_reference="ARR-009",
                              start_date="2025-01-01", end_date="2027-01-01")
    assert [v for v in _violations(await _export({_add: objs})) if v.startswith("806")] == []


async def test_a_parent_lei_that_is_a_subcontractor_keeps_the_subcontractor_data():
    rows = _sheets(await _export({_set: (Vendor, ("PP-004",), "ultimate_parent_id", "254900OPPU84GM83MG36")}))
    row = [r for r in rows("B_05.01") if r[0] == "254900OPPU84GM83MG36"]
    assert len(row) == 1 and row[0][4] == "DC Ops" and row[0][7]


async def test_two_subcontractor_records_with_one_lei_give_one_b05_02_row():
    def objs(pid):
        yield DoraSubcontractor(project_id=pid, id="SUB-009", lei="254900oppu84gm83mg36", name="DC Ops (duplicate)",
                                country_iso2="IE", person_type=_code("person_type"))
        yield DoraArrangementSubcontractor(project_id=pid, arrangement_id="ARR-001", subcontractor_id="SUB-009",
                                           tier=2, service_provided="S_01")
    assert [v for v in _violations(await _export({_add: objs})) if v.startswith("806")] == []


async def test_a_reporting_entity_also_in_the_consolidation_is_one_b01_02_row():
    from src.models import DoraConsolidationScope
    def objs(pid):
        yield DoraConsolidationScope(project_id=pid, id="CS-002", entity_lei=LEI_FE, entity_name="MedSecure SAS",
                                     country_iso2="FR", entity_type=_code("entity_type"), hierarchy="parent")
    assert [v for v in _violations(await _export({_add: objs})) if "B_01.02" in v] == []


async def test_the_server_names_the_subcontractor_served_by_its_name():
    from src.routes.dora import update_subcontractor_link
    from src.schemas import DoraArrangementSubcontractorUpdate
    async with _Session() as db:
        pid = uuid.uuid4()
        await _seed(db, pid)
        with pytest.raises(HTTPException) as exc:
            await update_subcontractor_link(pid, "ARR-001", "SUB-001",
                                            DoraArrangementSubcontractorUpdate(tier=3), user=None, db=db)
        assert "Power Ops" in exc.value.detail


async def test_a_provider_signer_known_as_a_vendor_keeps_the_vendor_data():
    from src.models import DoraSigner
    def objs(pid):
        yield Vendor(project_id=pid, id="PP-010", name="Cloud Corp Ireland", lei=LEI_OUTSIDE,
                     person_type=_code("person_type"), country_iso2="IE")
        yield DoraSigner(project_id=pid, arrangement_id="ARR-001", id="SIG-IE", signer_lei=LEI_OUTSIDE,
                         signer_name="Cloud Corp Ireland", signer_role="tpp")
    blob = await _export({_add: objs})
    assert _violations(blob) == []
    row = [r for r in _sheets(blob)("B_05.01") if r[0] == LEI_OUTSIDE]
    assert len(row) == 1 and row[0][4] == "Cloud Corp Ireland" and row[0][7]


async def test_a_lei_typed_with_spaces_is_the_same_group_lei_for_the_server():
    from src.models import DoraConsolidationScope
    from src.routes.dora import update_arrangement
    from src.schemas import DoraArrangementUpdate
    async with _Session() as db:
        pid = uuid.uuid4()
        await _seed(db, pid)
        cs = await db.get(DoraConsolidationScope, (pid, "CS-001"))
        cs.entity_lei = " " + cs.entity_lei.lower() + " "
        await db.commit()
        # Group IT (PP-003) carries the consolidation LEI: still intra-group, the link is accepted.
        out = await update_arrangement(pid, "ARR-002", DoraArrangementUpdate(linked_arrangement_id="ARR-003"),
                                       user=None, db=db)
        assert (out["linked_arrangement_id"] if isinstance(out, dict) else out.linked_arrangement_id) == "ARR-003"


async def test_the_parent_of_a_vendor_signing_for_a_provider_is_reported_too():
    from src.models import DoraSigner
    def objs(pid):
        yield Vendor(project_id=pid, id="PP-011", name="Cloud Holding", lei="5493006MHB84DD0ZWV18",
                     person_type=_code("person_type"), country_iso2="IE")
        yield Vendor(project_id=pid, id="PP-010", name="Cloud Corp Ireland", lei=LEI_OUTSIDE, ultimate_parent_id="PP-011",
                     person_type=_code("person_type"), country_iso2="IE")
        yield DoraSigner(project_id=pid, arrangement_id="ARR-001", id="SIG-IE", signer_lei=LEI_OUTSIDE,
                         signer_name="Cloud Corp Ireland", signer_role="tpp")
    assert _violations(await _export({_add: objs})) == []
