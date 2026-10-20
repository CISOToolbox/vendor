"""FEAT-60 — in the RoI Excel export, a data cell left empty although the EBA
data model requires its column is red; an empty optional cell, or a filled
required one, keeps the plain look. The cover sheet explains the red."""
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", ".."))
os.environ.setdefault("DATABASE_URL", "sqlite+aiosqlite://")
os.environ.setdefault("JWT_SECRET", "test-secret")

from openpyxl import Workbook  # noqa: E402

from src.dora_export import _MISSING_FILL, _write_sheet  # noqa: E402

# B_01.02: 0010 LEI and 0020 name required; 0100 currency and 0110 total assets optional.
COLS = [("0010", "LEI"), ("0020", "Name"), ("0100", "Currency"), ("0110", "Total assets")]


def _red(cell) -> bool:
    return cell.fill is not None and cell.fill.fgColor.rgb == _MISSING_FILL.fgColor.rgb and cell.fill.fill_type == "solid"


def test_empty_required_cells_are_red_and_the_rest_is_not():
    wb = Workbook()
    _write_sheet(wb, "B_01.02", COLS, [
        ["", "MedSecure SAS", "", None],          # LEI missing; optional columns empty
        ["213800WSGIIZCXF1P572", "   ", "EUR", 0],   # name blank; total assets 0 is a value
        [None, "Sub", "EUR", 10],                 # LEI missing
    ])
    ws = wb["B_01.02"]
    # Rows 1-2 are the header (titles, codes): never red.
    assert not any(_red(c) for c in ws[1]) and not any(_red(c) for c in ws[2])
    assert [_red(ws.cell(row=3, column=i)) for i in range(1, 5)] == [True, False, False, False]
    assert [_red(ws.cell(row=4, column=i)) for i in range(1, 5)] == [False, True, False, False]
    assert [_red(ws.cell(row=5, column=i)) for i in range(1, 5)] == [True, False, False, False]


def test_a_template_outside_the_model_is_left_plain():
    wb = Workbook()
    _write_sheet(wb, "B_99.01", [("0010", "id"), ("0030", "lei")], [["", ""]])
    ws = wb["B_99.01"]
    assert not any(_red(ws.cell(row=3, column=i)) for i in (1, 2))


def test_a_b02_02_sheet_marks_its_required_columns_too():
    # 0150 and 0160 required (0160 once misfiled under B_07.01), 0170 optional.
    from src.dora_export import _COLS_B0202
    codes = [c for c, _ in _COLS_B0202]
    row = ["x"] * len(codes)
    for c in ("0150", "0160", "0170"):
        row[codes.index(c)] = ""
    wb = Workbook()
    _write_sheet(wb, "B_02.02", _COLS_B0202, [row])
    ws = wb["B_02.02"]
    red = {codes[i - 1] for i in range(1, len(codes) + 1) if _red(ws.cell(row=3, column=i))}
    assert red == {"0150", "0160"}
