"""BUG-112 — a register filled in everything Vendor lets a user enter exports
with no required cell left empty (no red cell). Before the fix, ten required
columns could never be filled and eight could come out empty.

Builds the register in an in-memory SQLite (same shim as
test_dora_export_build.py), exports it and lists every red cell."""
import io
import json
import os
import uuid

import pytest

os.environ.setdefault("DATABASE_URL", "sqlite+aiosqlite://")
os.environ.setdefault("JWT_SECRET", "test-secret")

from openpyxl import load_workbook  # noqa: E402
from sqlalchemy import JSON  # noqa: E402
from sqlalchemy.dialects.postgresql import JSONB as _JSONB  # noqa: E402
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine  # noqa: E402
from sqlalchemy.pool import StaticPool  # noqa: E402

from src.dora_export import _MISSING_FILL, build_dora_xlsx  # noqa: E402
from src.models import (  # noqa: E402
    Base, DoraArrangement, DoraArrangementFunction, DoraArrangementRfe, DoraArrangementService,
    DoraArrangementSubcontractor, DoraBranch, DoraConsolidationScope, DoraEntity, DoraFunction,
    DoraSigner, DoraSubcontractor, Project, Vendor,
)

for _t in Base.metadata.tables.values():
    for _c in _t.columns:
        if _c.server_default is not None:
            _sd = str(getattr(_c.server_default, "arg", "")).lower()
            if any(k in _sd for k in ("gen_random_uuid", "now(", "::jsonb")):
                _c.server_default = None
        if isinstance(_c.type, _JSONB):
            _c.type = JSON()

_engine = create_async_engine("sqlite+aiosqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
_Session = async_sessionmaker(_engine, class_=AsyncSession, expire_on_commit=False)
pytestmark = pytest.mark.asyncio

_CL = json.load(open(os.path.join(os.path.dirname(__file__), "..", "..", "src", "dora_codelists.json")))


def _code(key: str) -> str:
    """First code of a codelist that the EBA knows (has an eba_code)."""
    return next(i["code"] for i in _CL[key] if i.get("eba_code"))


LEI_FE, LEI_SUB_FE, LEI_TPP, LEI_PARENT, LEI_INTRA, LEI_SUBC = (
    "529900T8BM49AURSDO55", "5493001KJTIIGC8Y1R12", "213800WSGIIZCXF1P572",
    "549300MLUDYVRQOOXS22", "5493001KJTIIGC8Y1R12", "254900OPPU84GM83MG36")


async def _seed(db, pid, extra=None):
    db.add(Project(id=pid, name="Complete register"))
    db.add(DoraEntity(project_id=pid, id="ENT-001", lei=LEI_FE, name="MedSecure SAS", country_iso2="FR",
                      competent_authority="ACPR", entity_type=_code("entity_type"), hierarchy="parent",
                      total_assets=1000.0, total_assets_currency="EUR", reporting_period="2026-12-31"))
    db.add(DoraBranch(project_id=pid, id="BR-001", rfe_id="ENT-001", name="MedSecure Lisboa", country_iso2="PT", branch_code="BR-PT"))
    db.add(DoraConsolidationScope(project_id=pid, id="CS-001", entity_lei=LEI_SUB_FE, entity_name="MedSecure Services",
                                  country_iso2="FR", entity_type=_code("entity_type"), hierarchy="subsidiary"))
    db.add(DoraFunction(project_id=pid, id="FN-001", code="F-PAY", name="Payments", is_critical_or_important=True,
                        criticality_rationale="core", business_line=_code("licenced_activity"),
                        recovery_time_objective_h=4, recovery_point_objective_h=1,
                        impact_tolerance_description=_code("impact_level"), last_assessment_date="2026-06-30",
                        rfe_id="ENT-001"))
    common = dict(person_type=_code("person_type"), country_iso2="IE")
    db.add(Vendor(project_id=pid, id="PP-001", name="Cloud Corp Holding", lei=LEI_PARENT, **common))
    db.add(Vendor(project_id=pid, id="PP-002", name="Cloud Corp", lei=LEI_TPP, ultimate_parent_id="PP-001", **common))
    db.add(Vendor(project_id=pid, id="PP-003", name="Group IT", lei=LEI_INTRA, **common))   # intra-group provider
    db.add(Vendor(project_id=pid, id="PP-004", name="No-LEI SaaS", additional_id_type="EUID",
                  additional_id_value="FRRCS.123456789", **common))                        # identified without a LEI
    arr = dict(arrangement_type=_code("arrangement_type"), currency="EUR", annual_cost_amount=1200.0,
               start_date="2025-01-01", jurisdiction_country="IE", governing_law_country="IE",
               data_processing_country="IE", data_sensitivity=_code("data_sensitivity"),
               reliance_level=_code("reliance_level"), substitutability_level=_code("substitutability"),
               reintegration_level=_code("reintegration_level"), impact_discontinuing_level=_code("impact_level"),
               last_audit_date="2026-03-31", exit_strategy_documented=True)
    db.add(DoraArrangement(project_id=pid, id="ARR-001", vendor_id="PP-002", arrangement_reference="ARR-001",
                           end_date="2028-12-31", data_stored=True, data_storage_country="IE", **arr))
    db.add(DoraArrangement(project_id=pid, id="ARR-002", vendor_id="PP-003", arrangement_reference="ARR-002",
                           open_ended=True, data_stored=False, linked_arrangement_id="ARR-001", **arr))
    db.add(DoraArrangement(project_id=pid, id="ARR-003", vendor_id="PP-004", arrangement_reference="ARR-003",
                           end_date="2027-06-30", data_stored=False, **arr))
    db.add(DoraSigner(project_id=pid, arrangement_id="ARR-003", id="SIG-TPP-3", signer_lei="", signer_name="SaaS CEO", signer_role="tpp"))
    # The provider of ARR-001 has a LEI and signs for itself: its row is already there.
    db.add(DoraSigner(project_id=pid, arrangement_id="ARR-001", id="SIG-TPP-1", signer_lei="", signer_name="Cloud CEO", signer_role="tpp"))
    for a in ("ARR-001", "ARR-002", "ARR-003"):
        db.add(DoraArrangementService(project_id=pid, arrangement_id=a, service_code="S_01"))
        db.add(DoraArrangementFunction(project_id=pid, arrangement_id=a, function_id="FN-001"))
        db.add(DoraArrangementRfe(project_id=pid, arrangement_id=a, rfe_id="ENT-001"))
        db.add(DoraSigner(project_id=pid, arrangement_id=a, id="SIG-" + a, signer_lei=LEI_FE, signer_name="CISO", signer_role="fe"))
    db.add(DoraSubcontractor(project_id=pid, id="SUB-001", lei=LEI_SUBC, name="DC Ops", country_iso2="IE",
                                  person_type=_code("person_type")))
    db.add(DoraArrangementSubcontractor(project_id=pid, arrangement_id="ARR-001", subcontractor_id="SUB-001",
                                        tier=2, service_provided="S_01"))
    db.add(DoraSubcontractor(project_id=pid, id="SUB-002", lei="5493000IBP32UQZ0KL24", name="Power Ops", country_iso2="IE",
                                  person_type=_code("person_type")))
    db.add(DoraArrangementSubcontractor(project_id=pid, arrangement_id="ARR-001", subcontractor_id="SUB-002",
                                        tier=3, service_provided="S_01", parent_subcontractor_id="SUB-001"))
    for k, v in (extra or {}).items():
        await k(db, pid, v)
    await db.commit()


def _red_cells(blob: bytes) -> list[str]:
    wb = load_workbook(io.BytesIO(blob))
    red = []
    for ws in wb.worksheets[1:]:
        for row in ws.iter_rows(min_row=3):
            for c in row:
                if c.fill is not None and c.fill.fill_type == "solid" and c.fill.fgColor.rgb == _MISSING_FILL.fgColor.rgb:
                    red.append(f"{ws.title}.{ws.cell(row=2, column=c.column).value} (row {c.row})")
    return red


async def _export(extra=None) -> bytes:
    pid = uuid.uuid4()
    async with _Session() as db:
        await _seed(db, pid, extra)
        await db.commit()
        proj = await db.get(Project, pid)
        return await build_dora_xlsx(db, pid, proj, target_currency="EUR")


async def _set(db, pid, change):
    model, key, field, value = change
    obj = await db.get(model, (pid,) + key)
    setattr(obj, field, value)


async def test_a_complete_register_has_no_required_cell_left_empty():
    async with _engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    try:
        blob = await _export()
        assert _red_cells(blob) == []
        wb = load_workbook(io.BytesIO(blob))
        rows = lambda name: [[c.value for c in r] for r in wb[name].iter_rows(min_row=3)]   # noqa: E731
        assert any(r[0] == "ARR-002" and r[1] == "ARR-001" for r in rows("B_02.03"))           # intra-group, linked
        assert any(r[0] == "ARR-002" and r[1] == LEI_INTRA for r in rows("B_03.03"))
        b0501 = {r[0]: r for r in rows("B_05.01")}
        assert b0501[LEI_TPP][10] == LEI_PARENT        # parent picked among vendors
        assert b0501[LEI_INTRA][10] == LEI_INTRA       # no group: its own code (FAQ Q99)
        b0202 = {r[0]: r for r in rows("B_02.02")}
        assert b0202["ARR-002"][7] == "9999-12-31"     # open-ended
        assert b0202["ARR-002"][14] == "eba_GA:qx2007" # nothing stored (FAQ Q61)
        assert all(r[3] == "Not Applicable" for r in rows("B_04.01"))   # not a branch (FAQ Q64)
        assert all(r[8] == "9999-12-31" for r in rows("B_01.02"))       # in scope, not deleted
        b0302 = [(r[0], r[1]) for r in rows("B_03.02")]
        assert len(b0302) == len(set(b0302))                            # one key, one row
        assert ("ARR-003", "FRRCS.123456789") in b0302                  # no LEI: its EUID, not the internal id
        assert not [k for k in b0302 if str(k[1]).startswith("PP-")]    # never an internal id
        assert b0302.count(("ARR-001", LEI_TPP)) == 1                   # a provider signing for itself: one row
        assert LEI_PARENT in b0501                                      # the cited parent has a row of its own
        assert all(not str(r[0]).startswith("PP-") for r in rows("B_05.01"))
        b0502 = [r for r in rows("B_05.02") if r[4] == 3]
        assert b0502 and b0502[0][5] == LEI_SUBC                        # rank 3: recipient = rank 2 sub
    finally:
        async with _engine.begin() as conn:
            await conn.run_sync(Base.metadata.drop_all)


@pytest.mark.parametrize("change,red", [
    ((DoraArrangement, ("ARR-001",), "exit_strategy_documented", None), "B_07.01.0080"),
    ((DoraArrangementSubcontractor, ("ARR-001", "SUB-002"), "parent_subcontractor_id", None), "B_05.02.0060"),
    ((Vendor, ("PP-002",), "ultimate_parent_id", "PP-GONE"), "B_05.01.0110"),        # orphan, not a LEI
    ((DoraArrangement, ("ARR-002",), "data_stored", None), "B_02.02.0140"),            # not stated, no country: unknown, never a silent "No"
    ((Vendor, ("PP-004",), "additional_id_value", ""), "B_05.01.0010"),               # no code at all
    ((DoraSubcontractor, ("SUB-001",), "person_type", ""), "B_05.01.0070"),            # a subcontractor is a B_05.01 row too
])
async def test_an_unstated_value_shows_red(change, red):
    async with _engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    try:
        blob = await _export({_set: change})
        assert [c for c in _red_cells(blob) if c.startswith(red)], _red_cells(blob)
    finally:
        async with _engine.begin() as conn:
            await conn.run_sync(Base.metadata.drop_all)


async def test_free_text_is_refused_as_function_impact():
    from src.dora_validation import validate_dora_function
    with pytest.raises(Exception) as exc:
        validate_dora_function({"impact_tolerance_description": "the business stops"})
    assert exc.value.status_code == 422


async def test_the_parent_is_saved_through_the_route_and_cleared_when_deleted():
    from src.routes.dora import patch_vendor_roi
    from src.routes.vendors import delete_vendor
    from src.schemas import DoraVendorRoIPatch
    async with _engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    try:
        pid = uuid.uuid4()
        async with _Session() as db:
            await _seed(db, pid)
            await db.commit()
            out = await patch_vendor_roi(pid, "PP-004", DoraVendorRoIPatch(ultimate_parent_id="PP-001"), user=None, db=db)
            assert out["ultimate_parent_id"] == "PP-001"                    # was refused (422) as "not a LEI"
            await delete_vendor(pid, "PP-001", user=None, db=db)
            for vid in ("PP-002", "PP-004"):
                v = await db.get(Vendor, (pid, vid))
                await db.refresh(v)
                assert v.ultimate_parent_id is None                         # no orphan parent left
    finally:
        async with _engine.begin() as conn:
            await conn.run_sync(Base.metadata.drop_all)


async def _rows_after(extra):
    blob = await _export(extra)
    wb = load_workbook(io.BytesIO(blob))
    return lambda name: [[c.value for c in r] for r in wb[name].iter_rows(min_row=3)]   # noqa: E731


async def test_a_provider_without_any_code_gives_one_b03_02_row_per_arrangement():
    async with _engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    try:
        rows = await _rows_after({_set: (Vendor, ("PP-004",), "additional_id_value", "")})
        assert len([r for r in rows("B_03.02") if r[0] == "ARR-003"]) == 1   # empty key: one row, not two
    finally:
        async with _engine.begin() as conn:
            await conn.run_sync(Base.metadata.drop_all)


async def test_a_former_lei_parent_is_exported_in_upper_case():
    async with _engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    try:
        rows = await _rows_after({_set: (Vendor, ("PP-004",), "ultimate_parent_id", LEI_PARENT.lower())})
        b0501 = [r for r in rows("B_05.01") if r[2] == "FRRCS.123456789" or r[0] == "FRRCS.123456789"]
        assert b0501 and b0501[0][10] == LEI_PARENT
    finally:
        async with _engine.begin() as conn:
            await conn.run_sync(Base.metadata.drop_all)


async def test_removing_a_vendor_clears_the_arrangements_linking_to_its_own():
    from src.routes.vendors import delete_vendor
    async with _engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    try:
        pid = uuid.uuid4()
        async with _Session() as db:
            await _seed(db, pid)
            await db.commit()
            await delete_vendor(pid, "PP-004", user=None, db=db)          # unrelated: the link stays
            a2 = await db.get(DoraArrangement, (pid, "ARR-002"))
            await db.refresh(a2)
            assert a2.linked_arrangement_id == "ARR-001"
            await delete_vendor(pid, "PP-002", user=None, db=db)          # vendor of ARR-001
            await db.refresh(a2)
            assert a2.linked_arrangement_id is None
    finally:
        async with _engine.begin() as conn:
            await conn.run_sync(Base.metadata.drop_all)
